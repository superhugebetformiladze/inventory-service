"""План и факт выпуска продукции: разбор рапорта Excel и сводки по нему.

Рапорт — это «Справка о выполнении плана производства заготовительными цехами».
Устроена она так:

    Справка о выполнении плана производства заготовительными цехами
    Склад отправитель        | План | Сдано на склад ПДО,шт. | Возврат | Примечание
    Номенклатура | Тип корпуса
    Цех №2 - Инструментально-штамповочный | 198341 | 194826          <- итог по цеху
    Изделие А    | Корпус А  | 700    | 700                          <- строки цеха
    Изделие Б    | Корпус Б  | 30000  | 30000
    Цех №3 - Сборки корпусов              | 84226  | 76176           <- следующий цех
    ...
    Итого                                 | 311776 | 297842

Важные особенности, из которых вытекает устройство модуля:

  * **станков в рапорте нет** — он составляется по цехам. Поэтому связь «станок —
    изделие» заводится вручную в админке (таблица MachineProduct), а не берётся
    из файла;
  * **периода в файле тоже нет** — за какой месяц справка, знает только тот, кто её
    выгрузил. Поэтому период указывается при загрузке;
  * строки цехов и строка «Итого» — это суммы, их нельзя складывать со строками
    номенклатуры. Мы их не сохраняем, но используем для проверки: если сумма
    разобранных строк не сошлась с итогом из файла, об этом сообщается в отчёте.
"""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import MachineProduct, Product, ProductionReport, ProductionRow

# Заголовки колонок рапорта. Ищем по вхождению, поэтому «Сдано на склад ПДО,шт.»
# найдётся и в варианте «Сдано на склад ПДО, шт».
COLUMN_SYNONYMS: dict[str, tuple[str, ...]] = {
    "name": ("номенклатура", "изделие", "наименование", "продукция", "склад отправитель"),
    "body_type": ("тип корпуса", "тип изделия", "корпус"),
    "plan": ("план",),
    "fact": ("сдано на склад", "сдано", "факт", "выпущено", "изготовлено"),
    "return_qty": ("возврат",),
    "note": ("примечание", "комментарий"),
}

# «Цех №2 - Инструментально-штамповочный» — строка-итог по цеху, не номенклатура
SHOP_RE = re.compile(r"^\s*цех\b", re.IGNORECASE)
TOTAL_RE = re.compile(r"^\s*(итого|всего|итог)\b", re.IGNORECASE)
# из «Цех №2 - Инструментально-штамповочный» достаём номер цеха
SHOP_NUM_RE = re.compile(r"№\s*(\d+)")

MAX_HEADER_SCAN = 20  # шапка рапорта всегда в начале листа


class ReportParseError(ValueError):
    """Файл не похож на рапорт — сообщаем человеку, чего именно не хватило."""


@dataclass
class ParsedRow:
    """Одна строка номенклатуры из файла."""

    row_no: int
    name: str
    body_type: str | None
    workshop: str | None
    plan_qty: float | None
    fact_qty: float | None
    return_qty: float | None
    note: str | None


@dataclass
class ParsedReport:
    rows: list[ParsedRow] = field(default_factory=list)
    sheet: str = ""
    # Итоги, посчитанные самим файлом, — для сверки с суммой разобранных строк
    file_plan_total: float | None = None
    file_fact_total: float | None = None
    shops: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def plan_total(self) -> float:
        return sum(row.plan_qty or 0 for row in self.rows)

    @property
    def fact_total(self) -> float:
        return sum(row.fact_qty or 0 for row in self.rows)


def _text(value: Any) -> str:
    """Значение ячейки как аккуратная строка; «.» и пустые считаются пустыми."""
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return "" if text in {".", "-", "—", "–"} else text


def _normalize(value: Any) -> str:
    return _text(value).lower()


def to_number(value: Any) -> float | None:
    """Число из ячейки. Пустые клетки и заглушки вида «.» дают None, а не ноль.

    Разница важна: «плана не было» и «план равен нулю» — это разные вещи,
    и в проценте выполнения они ведут себя по-разному.
    """
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = _text(value).replace(" ", "").replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def product_key(name: str, body_type: str | None, workshop: str | None) -> str:
    """Ключ номенклатуры: по нему повторная загрузка находит уже заведённое изделие."""
    parts = [_normalize(name), _normalize(body_type), _normalize(workshop)]
    return "|".join(parts)[:600]


def shop_number(workshop: str | None) -> str | None:
    """«Цех №2 - Инструментально-штамповочный» -> «2»."""
    if not workshop:
        return None
    found = SHOP_NUM_RE.search(workshop)
    return found.group(1) if found else None


def _find_columns(rows: list[list[Any]]) -> tuple[dict[str, int], int]:
    """Находит, в каких колонках лежат название, план и факт, и где кончается шапка.

    Шапка двухуровневая: «План» и «Сдано на склад ПДО» стоят строкой выше, чем
    «Номенклатура» и «Тип корпуса». Поэтому ищем строку со словом «План», берём её
    вместе со следующей и склеиваем заголовки по колонкам.
    """
    header_index = None
    for index, row in enumerate(rows[:MAX_HEADER_SCAN]):
        if any(_normalize(cell) == "план" for cell in row):
            header_index = index
            break
    if header_index is None:
        raise ReportParseError(
            "В файле не нашлась шапка рапорта: нужна колонка «План» "
            "и колонка со сданным количеством"
        )

    header_rows = [rows[header_index]]
    following = rows[header_index + 1] if header_index + 1 < len(rows) else []
    # Вторая строка шапки — та, где только подписи и ни одного числа
    if following and not any(to_number(cell) is not None for cell in following):
        header_rows.append(following)

    width = max((len(row) for row in header_rows), default=0)
    combined: list[str] = []
    for column in range(width):
        texts = [_normalize(row[column]) for row in header_rows if column < len(row)]
        combined.append(" ".join(text for text in texts if text))

    columns: dict[str, int] = {}
    for column, title in enumerate(combined):
        if not title:
            continue
        for field_name, synonyms in COLUMN_SYNONYMS.items():
            if field_name in columns:
                continue
            if any(synonym in title for synonym in synonyms):
                columns[field_name] = column
                break

    if "plan" not in columns or "fact" not in columns:
        raise ReportParseError(
            "В шапке рапорта не нашлись колонки «План» и «Сдано на склад ПДО»"
        )
    columns.setdefault("name", 0)
    return columns, header_index + len(header_rows)


def parse_report(content: bytes) -> ParsedReport:
    """Читает книгу Excel и достаёт строки номенклатуры с планом и фактом."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - человеку нужна причина, а не трассировка
        raise ReportParseError("Не удалось прочитать файл Excel: {0}".format(exc)) from exc

    sheet = workbook.worksheets[0]
    raw = [list(row) for row in sheet.iter_rows(values_only=True)]
    return parse_grid(raw, sheet.title)


def parse_grid(raw: list[list[Any]], sheet_title: str = "") -> ParsedReport:
    """Разбор уже прочитанного листа. Вынесено отдельно, чтобы логику можно было
    проверить на тестовых данных, не поднимая Excel."""
    if not raw:
        raise ReportParseError("Лист рапорта пуст")

    columns, start = _find_columns(raw)
    report = ParsedReport(sheet=sheet_title)

    def cell(row: list[Any], key: str) -> Any:
        index = columns.get(key)
        return row[index] if index is not None and index < len(row) else None

    workshop: str | None = None
    for offset, row in enumerate(raw[start:], start=start + 1):
        name = _text(cell(row, "name"))
        plan = to_number(cell(row, "plan"))
        fact = to_number(cell(row, "fact"))

        if not name and plan is None and fact is None:
            continue

        # Строка «Итого» — общий итог файла, он нужен только для сверки
        if TOTAL_RE.match(name):
            report.file_plan_total = plan
            report.file_fact_total = fact
            continue

        # Строка цеха — заголовок блока с итогами по цеху; сами числа не сохраняем
        if SHOP_RE.match(name):
            workshop = name
            report.shops.append(name)
            continue

        if not name:
            report.warnings.append("строка {0}: есть числа, но нет названия".format(offset))
            continue

        report.rows.append(
            ParsedRow(
                row_no=offset,
                name=name[:255],
                body_type=(_text(cell(row, "body_type")) or None),
                workshop=workshop,
                plan_qty=plan,
                fact_qty=fact,
                return_qty=to_number(cell(row, "return_qty")),
                note=(_text(cell(row, "note")) or None),
            )
        )

    if not report.rows:
        raise ReportParseError("В файле не нашлось ни одной строки номенклатуры")
    return report


@dataclass
class ImportReport:
    """Что получилось при загрузке рапорта — показывается в админке."""

    rows: int = 0
    products_created: int = 0
    products_matched: int = 0
    shops: list[str] = field(default_factory=list)
    plan_total: float = 0.0
    fact_total: float = 0.0
    checks: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "rows": self.rows,
            "products_created": self.products_created,
            "products_matched": self.products_matched,
            "shops": self.shops,
            "plan_total": self.plan_total,
            "fact_total": self.fact_total,
            "checks": self.checks,
            "warnings": self.warnings,
        }


def save_upload(content: bytes, suffix: str = ".xlsx") -> str:
    """Кладёт исходный файл рядом с XML-выгрузками под сгенерированным именем."""
    stored_name = "{0}{1}".format(uuid.uuid4().hex, suffix)
    (settings.storage_dir / stored_name).write_bytes(content)
    return stored_name


def find_period(db: Session, period_from: date, period_to: date) -> ProductionReport | None:
    """Уже загруженный рапорт за тот же период — чтобы не задваивать план и факт."""
    return db.scalar(
        select(ProductionReport).where(
            ProductionReport.period_from == period_from,
            ProductionReport.period_to == period_to,
        )
    )


def get_or_create_product(
    db: Session,
    cache: dict[str, Product],
    name: str,
    body_type: str | None,
    workshop: str | None,
    counters: ImportReport | None = None,
) -> Product:
    key = product_key(name, body_type, workshop)
    found = cache.get(key)
    if found is not None:
        if counters is not None:
            counters.products_matched += 1
        return found

    product = Product(
        match_key=key,
        name=name,
        body_type=body_type,
        workshop=workshop,
    )
    db.add(product)
    db.flush()
    cache[key] = product
    if counters is not None:
        counters.products_created += 1
    return product


def import_report(
    db: Session,
    content: bytes,
    original_name: str,
    period_from: date,
    period_to: date,
    uploaded_by_id: int | None = None,
    note: str | None = None,
) -> tuple[ProductionReport, ImportReport]:
    """Разбирает рапорт и сохраняет план и факт за указанный период."""
    parsed = parse_report(content)
    summary = ImportReport(shops=parsed.shops, warnings=parsed.warnings)

    stored_name = save_upload(content)
    report = ProductionReport(
        original_name=original_name[:512],
        stored_name=stored_name,
        size_bytes=len(content),
        period_from=period_from,
        period_to=period_to,
        uploaded_by_id=uploaded_by_id,
        note=note,
    )
    db.add(report)
    db.flush()

    cache: dict[str, Product] = {
        product.match_key: product for product in db.scalars(select(Product))
    }

    for row in parsed.rows:
        product = get_or_create_product(
            db, cache, row.name, row.body_type, row.workshop, summary
        )
        db.add(
            ProductionRow(
                report_id=report.id,
                product_id=product.id,
                period_from=period_from,
                period_to=period_to,
                plan_qty=row.plan_qty,
                fact_qty=row.fact_qty,
                return_qty=row.return_qty,
                workshop=row.workshop,
                note=row.note,
                row_no=row.row_no,
                source="file",
            )
        )

    summary.rows = len(parsed.rows)
    summary.plan_total = parsed.plan_total
    summary.fact_total = parsed.fact_total

    # Сверка с итогом самого файла: если не сошлось, значит часть строк не разобралась
    for label, ours, theirs in (
        ("плану", parsed.plan_total, parsed.file_plan_total),
        ("факту", parsed.fact_total, parsed.file_fact_total),
    ):
        if theirs is None:
            continue
        if abs(ours - theirs) < 0.5:
            summary.checks.append("Итог по {0} сошёлся с файлом: {1:g}".format(label, theirs))
        else:
            summary.checks.append(
                "Итог по {0} разошёлся: в файле {1:g}, по разобранным строкам {2:g}".format(
                    label, theirs, ours
                )
            )

    report.rows_count = summary.rows
    report.products_count = summary.products_created + summary.products_matched
    report.plan_total = summary.plan_total
    report.fact_total = summary.fact_total
    db.commit()
    db.refresh(report)
    return report, summary


def delete_report(db: Session, report: ProductionReport) -> None:
    """Удаляет рапорт вместе со строками; заведённая номенклатура остаётся."""
    db.execute(delete(ProductionRow).where(ProductionRow.report_id == report.id))
    if report.stored_name:
        path = settings.storage_dir / report.stored_name
        path.unlink(missing_ok=True)
    db.delete(report)
    db.commit()


# --------------------------------------------------------------------------------------
# Сводки
# --------------------------------------------------------------------------------------


def _percent(fact: float | None, plan: float | None) -> float | None:
    """Процент выполнения = факт ÷ план. Без плана процента не существует."""
    if not plan:
        return None
    return round((fact or 0) / plan * 100, 1)


def rows_in_period(
    db: Session,
    date_from: date | None = None,
    date_to: date | None = None,
    product_ids: list[int] | None = None,
    machine_key: str | None = None,
):
    """Строки плана и факта, пересекающиеся с запрошенным периодом.

    Пересечение, а не строгое вхождение: месячный рапорт должен попадать в отчёт
    и тогда, когда пользователь смотрит неделю внутри этого месяца.
    """
    stmt = select(ProductionRow)
    if date_from:
        stmt = stmt.where(ProductionRow.period_to >= date_from)
    if date_to:
        stmt = stmt.where(ProductionRow.period_from <= date_to)
    if product_ids is not None:
        if not product_ids:
            return []
        stmt = stmt.where(ProductionRow.product_id.in_(product_ids))
    if machine_key:
        stmt = stmt.where(ProductionRow.machine_key == machine_key)
    return list(db.scalars(stmt.order_by(ProductionRow.period_from, ProductionRow.row_no)))


def machine_product_ids(db: Session, machine_key: str) -> list[int]:
    return list(
        db.scalars(select(MachineProduct.product_id).where(MachineProduct.machine_key == machine_key))
    )


def aggregate(rows: list[ProductionRow]) -> dict[str, Any]:
    """Свод по списку строк: план, факт, возврат и процент выполнения."""
    plan = sum(row.plan_qty or 0 for row in rows)
    fact = sum(row.fact_qty or 0 for row in rows)
    returns = sum(row.return_qty or 0 for row in rows)
    return {
        "plan_qty": plan,
        "fact_qty": fact,
        "return_qty": returns,
        "percent": _percent(fact, plan),
        "rows": len(rows),
    }


def by_product(rows: list[ProductionRow]) -> list[dict[str, Any]]:
    """Свод по каждому изделию: план, факт, процент."""
    buckets: dict[int, dict[str, Any]] = {}
    for row in rows:
        item = buckets.setdefault(
            row.product_id,
            {
                "product_id": row.product_id,
                "name": row.product.name,
                "body_type": row.product.body_type,
                "workshop": row.workshop or row.product.workshop,
                "unit": row.product.unit,
                "plan_qty": 0.0,
                "fact_qty": 0.0,
                "return_qty": 0.0,
                "rows": 0,
                "manual": False,
            },
        )
        item["plan_qty"] += row.plan_qty or 0
        item["fact_qty"] += row.fact_qty or 0
        item["return_qty"] += row.return_qty or 0
        item["rows"] += 1
        if row.source == "manual":
            item["manual"] = True

    result = list(buckets.values())
    for item in result:
        item["percent"] = _percent(item["fact_qty"], item["plan_qty"])
    result.sort(key=lambda item: (-(item["plan_qty"] or 0), item["name"]))
    return result


def machine_production(
    db: Session,
    machine_key: str,
    date_from: date | None = None,
    date_to: date | None = None,
) -> dict[str, Any]:
    """Продукция станка за период: что за ним закреплено и как выполнен план.

    Рапорт цеховой, поэтому у строк из файла станка нет. Для станка берём строки
    по закреплённой за ним номенклатуре и отдельно помечаем, сколько станков делят
    между собой каждое изделие — иначе цифру легко принять за выпуск одного станка.
    """
    links = list(
        db.scalars(select(MachineProduct).where(MachineProduct.machine_key == machine_key))
    )
    product_ids = [link.product_id for link in links]

    # Строки по закреплённой номенклатуре (цеховые) и строки, заведённые прямо на станок
    shared = rows_in_period(db, date_from, date_to, product_ids=product_ids) if product_ids else []
    own = rows_in_period(db, date_from, date_to, machine_key=machine_key)
    own_ids = {row.id for row in own}
    combined = own + [row for row in shared if row.id not in own_ids]

    # Сколько станков закреплено за каждым изделием — чтобы честно подписать цифры
    sharers: dict[int, int] = {}
    if product_ids:
        counted = db.execute(
            select(MachineProduct.product_id, func.count(MachineProduct.id))
            .where(MachineProduct.product_id.in_(product_ids))
            .group_by(MachineProduct.product_id)
        )
        sharers = {product_id: count for product_id, count in counted}

    products = by_product(combined)
    for item in products:
        item["machines"] = sharers.get(item["product_id"], 1)

    return {
        "machine_key": machine_key,
        "totals": aggregate(combined),
        "products": products,
        "linked": len(links),
        "shared": any(item.get("machines", 1) > 1 for item in products),
    }


# --------------------------------------------------------------------------------------
# Выгрузка в Excel
# --------------------------------------------------------------------------------------

# Excel исполняет ячейку, начинающуюся с = + - @ как формулу. Названия изделий приходят
# из чужого файла, поэтому такие значения помечаем апострофом — Excel покажет их текстом.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")


def _safe_text(value: Any) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in FORMULA_START else text


def _machine_names(db: Session) -> dict[str, str]:
    """machine_key -> человеческое название станка из карточки оборудования."""
    from . import equipment

    return equipment.titles(db)


def build_workbook(
    db: Session,
    rows: list[ProductionRow],
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    """Книга Excel с планом и фактом: строки, свод по цехам и свод по изделиям.

    Числа кладём числами, а не текстом, чтобы в Excel по ним сразу считались
    суммы и строились сводные таблицы.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    titles = _machine_names(db)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")
    total_font = Font(bold=True)
    INT = "# ##0"
    PCT = "0.0"
    DATE = "DD.MM.YYYY"

    workbook = Workbook()

    def write_sheet(title: str, headers: list[tuple[str, int]], data: list[list[Any]],
                    totals: list[Any] | None = None, number_from: int = 0) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append([name for name, _ in headers])
        for index, (_, width) in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 28
        sheet.freeze_panes = "A2"

        for line in data:
            sheet.append(line)
        last_data_row = sheet.max_row

        for row_cells in sheet.iter_rows(min_row=2):
            for cell in row_cells:
                # datetime — подкласс date, поэтому одной проверки хватает на оба типа
                if isinstance(cell.value, date):
                    cell.number_format = DATE
                elif isinstance(cell.value, (int, float)) and cell.column > number_from:
                    cell.number_format = PCT if "%" in headers[cell.column - 1][0] else INT

        if totals is not None:
            sheet.append(totals)
            for cell in sheet[sheet.max_row]:
                cell.font = total_font
                if isinstance(cell.value, (int, float)):
                    cell.number_format = PCT if "%" in headers[cell.column - 1][0] else INT

        # Фильтр не должен захватывать строку «Итого»: иначе сортировка в Excel
        # утащит её в середину таблицы
        if last_data_row > 1:
            sheet.auto_filter.ref = "A1:{0}{1}".format(
                get_column_letter(len(headers)), last_data_row
            )

    # --- лист со строками ---
    detail_headers = [
        ("Изделие", 42), ("Тип корпуса", 26), ("Цех", 34), ("Станок", 26),
        ("Период с", 13), ("Период по", 13), ("План", 12), ("Факт", 12),
        ("Выполнение, %", 14), ("Возврат", 11), ("Источник", 12), ("Примечание", 32),
    ]
    detail: list[list[Any]] = []
    for row in sorted(rows, key=lambda item: (item.workshop or "", item.product.name)):
        plan = row.plan_qty
        fact = row.fact_qty
        detail.append([
            _safe_text(row.product.name),
            _safe_text(row.product.body_type),
            _safe_text(row.workshop or row.product.workshop),
            _safe_text(titles.get(row.machine_key or "", row.machine_key or "")),
            row.period_from,
            row.period_to,
            plan,
            fact,
            round((fact or 0) / plan * 100, 1) if plan else None,
            row.return_qty,
            "рапорт" if row.source == "file" else "вручную",
            _safe_text(row.note),
        ])
    summary = aggregate(rows)
    write_sheet(
        "План и факт",
        detail_headers,
        detail,
        totals=["Итого", "", "", "", "", "", summary["plan_qty"], summary["fact_qty"],
                summary["percent"], summary["return_qty"], "", ""],
        number_from=6,
    )

    # --- свод по цехам ---
    shops: dict[str, list[ProductionRow]] = {}
    for row in rows:
        shops.setdefault(row.workshop or row.product.workshop or "Без цеха", []).append(row)
    shop_rows = []
    for name, items in shops.items():
        shop = aggregate(items)
        shop_rows.append([
            _safe_text(name), shop["plan_qty"], shop["fact_qty"],
            shop["percent"], shop["return_qty"], shop["rows"],
        ])
    shop_rows.sort(key=lambda line: -(line[1] or 0))
    write_sheet(
        "Свод по цехам",
        [("Цех", 44), ("План", 14), ("Факт", 14), ("Выполнение, %", 14),
         ("Возврат", 12), ("Строк", 10)],
        shop_rows,
        totals=["Итого", summary["plan_qty"], summary["fact_qty"], summary["percent"],
                summary["return_qty"], summary["rows"]],
        number_from=1,
    )

    # --- свод по изделиям ---
    write_sheet(
        "Свод по изделиям",
        [("Изделие", 44), ("Тип корпуса", 26), ("Цех", 34), ("План", 13),
         ("Факт", 13), ("Выполнение, %", 14), ("Возврат", 11)],
        [[_safe_text(item["name"]), _safe_text(item["body_type"]), _safe_text(item["workshop"]),
          item["plan_qty"], item["fact_qty"], item["percent"], item["return_qty"]]
         for item in by_product(rows)],
        totals=["Итого", "", "", summary["plan_qty"], summary["fact_qty"],
                summary["percent"], summary["return_qty"]],
        number_from=3,
    )

    # --- титульный лист с параметрами выборки ---
    info = workbook["Sheet"]
    info.title = "О выгрузке"
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 46
    period = (
        "все загруженные данные"
        if not date_from and not date_to
        else "{0} — {1}".format(
            date_from.strftime("%d.%m.%Y") if date_from else "…",
            date_to.strftime("%d.%m.%Y") if date_to else "…",
        )
    )
    for label, value in (
        ("Отчёт", "План и факт выпуска продукции"),
        ("Период", period),
        ("Сформирован", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Строк плана и факта", summary["rows"]),
        ("План, шт.", summary["plan_qty"]),
        ("Сдано, шт.", summary["fact_qty"]),
        ("Выполнение, %", summary["percent"]),
        ("Возврат, шт.", summary["return_qty"]),
    ):
        info.append([label, value])
    for cell in info["A"]:
        cell.font = total_font

    workbook.move_sheet("О выгрузке", offset=3)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
